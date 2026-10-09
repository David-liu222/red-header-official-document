#!/usr/bin/env python3
"""Read-only structural validator for a red-header official-document DOCX.

This intentionally reports evidence and does not rewrite the document. It is a
first-pass check; red-header color/line placement and visual balance still need
rendered page-image inspection.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from datetime import date as calendar_date
from pathlib import Path

from apply_red_header_docx import (
    DEFAULT_NUMBER_LINE_PT,
    RED_LINE_WIDTH_PT,
    SIGNER_LABEL_TRACKING_TWIPS,
    SIGNER_RIGHT_INDENT_PT,
    canonical_title,
    official_filename,
)
from red_header_common import (attachment_start, closing, independent_heading, numbered_level,
                               safe_blank, spacing_issues, indent_issues, run_format_issues,
                               effective_ppr, on)

try:
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.oxml.ns import qn
except ImportError as exc:  # pragma: no cover - environment-dependent
    raise SystemExit("缺少 python-docx，请在工作环境安装后重试。") from exc


def east_asia_font(run):
    """Return the East Asian font name when it is explicitly present."""
    rpr = run._element.rPr
    if rpr is None or rpr.rFonts is None:
        return None
    return rpr.rFonts.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia")


def mm(value):
    return round(value / Mm(1), 2)


def pt(value):
    return None if value is None else round(value / Pt(1), 2)


def nonempty_paragraphs(doc):
    return [(idx + 1, p) for idx, p in enumerate(doc.paragraphs) if p.text.strip()]


def is_signer_paragraph(paragraph):
    return bool(re.match(r"^签\s*发\s*人\s*[:：]", paragraph.text.strip()))


def has_paragraph_border(paragraph):
    return bool(paragraph._p.xpath("./w:pPr/w:pBdr"))


TABLE_RE = re.compile(rb"<w:tbl(?:\s[^>]*)?>.*?</w:tbl>", re.DOTALL)
TBL_BORDERS_RE = re.compile(rb"<w:tblBorders(?:\s[^>]*)?>(.*?)</w:tblBorders>", re.DOTALL)
TC_BORDERS_RE = re.compile(rb"<w:tcBorders(?:\s[^>]*)?>(.*?)</w:tcBorders>", re.DOTALL)
BORDER_ELEMENT_RE = re.compile(rb"<w:(top|left|bottom|right|insideH|insideV|tl2br|tr2bl)\b([^>]*)/>")
REQUIRED_TABLE_BORDERS = (b"top", b"left", b"bottom", b"right", b"insideH", b"insideV")


def attr_value(attrs: bytes, name: bytes) -> bytes | None:
    match = re.search(rb"\s" + re.escape(name) + rb'="([^"]*)"', attrs)
    return match.group(1) if match else None


def table_border_issues(docx_path: Path) -> tuple[int, list[dict]]:
    """Return issues for non-solid/missing borders and unsafe floating tables."""
    with zipfile.ZipFile(docx_path, "r") as package:
        document_xml = package.read("word/document.xml")
    issues: list[dict] = []
    tables = TABLE_RE.findall(document_xml)
    for table_idx, table_xml in enumerate(tables, 1):
        if b"<w:tblpPr" in table_xml:
            issues.append({"table": table_idx, "issue": "floating_table_text_wrapping", "tag": "w:tblpPr"})
        tbl_borders = TBL_BORDERS_RE.search(table_xml)
        if not tbl_borders:
            issues.append({"table": table_idx, "issue": "missing_tblBorders"})
        else:
            present = {name for name, attrs in BORDER_ELEMENT_RE.findall(tbl_borders.group(1))}
            for required in REQUIRED_TABLE_BORDERS:
                if required not in present:
                    issues.append({"table": table_idx, "issue": "missing_table_border", "border": required.decode()})
        border_scopes = []
        if tbl_borders:
            border_scopes.append(tbl_borders.group(1))
        border_scopes.extend(match.group(1) for match in TC_BORDERS_RE.finditer(table_xml))
        for scope in border_scopes:
            for name, attrs in BORDER_ELEMENT_RE.findall(scope):
                value = attr_value(attrs, b"w:val")
                if value != b"single":
                    issues.append({
                        "table": table_idx,
                        "border": name.decode(),
                        "val": None if value is None else value.decode(errors="replace"),
                        "issue": "not_single",
                    })
    return len(tables), issues


def main():
    parser = argparse.ArgumentParser(description="检查套红头 DOCX 的常见固定规则（只读）")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--title", help="正文标题，用于校验文件名和标题段落")
    parser.add_argument("--issuing-number", help="正式文号；省略时读取带红色下边框的文号段落")
    parser.add_argument("--expected-company", default="内蒙古准格尔旗力量煤业有限公司")
    parser.add_argument("--expected-mine", help="需要校验时传入矿名；不传则不强制要求矿名")
    parser.add_argument("--body-bold", action="store_true", help="仅用户明确要求普通正文加粗时使用；一级独立小标题默认加粗")
    parser.add_argument("--no-recipient", action="store_true", help="与生成器同步：明确原稿无主送，不把首段正文当主送")
    parser.add_argument("--require-manual-review", action="store_true", help="人工项未完成时返回2；默认返回码仅反映自动硬性不符合")
    parser.add_argument("--allow-table-spacing-adjustment", action="store_true", help="仅有真实表格且已记录局部压缩原因时，允许25—28磅正文；仍须渲染")
    parser.add_argument("--sign-type", choices=("external", "internal"), default="external")
    parser.add_argument("--inspection-material", action="store_true", help="迎检文件必须无签发人/印发栏")
    parser.add_argument("--blank-document-number", action="store_true", help="按用户要求核验红头后第3段文号正文留空")
    parser.add_argument("--blank-number-line-pt", type=float, default=DEFAULT_NUMBER_LINE_PT,
                        help="空文号段预期行距，默认27磅；本次2023内发样本使用19磅")
    parser.add_argument("--signer-label-tracking-twips", type=int, default=SIGNER_LABEL_TRACKING_TWIPS,
                        help="按同类型正式样本核验签发栏‘签发’字距，默认140")
    parser.add_argument("--signer-right-indent-pt", type=float, default=SIGNER_RIGHT_INDENT_PT,
                        help="按同类型正式样本核验签发栏右缩进，默认68磅")
    parser.add_argument("--json", action="store_true", help="输出 JSON")
    args = parser.parse_args()

    if args.docx.suffix.lower() != ".docx":
        raise SystemExit("只接受 .docx；旧 .doc 请先在副本上转换。")
    doc = Document(str(args.docx))
    checks = []

    def add(rule, status, evidence, severity="info"):
        if status == 'review' and severity == 'error':
            status = 'fail'
        checks.append({"rule": rule, "status": status, "evidence": evidence, "severity": severity})

    for idx, section in enumerate(doc.sections, 1):
        size_ok = abs(mm(section.page_width) - 210) < 0.2 and abs(mm(section.page_height) - 297) < 0.2
        add(f'第{idx}节A4纸张', 'pass' if size_ok else 'fail',
            {'width_mm': mm(section.page_width), 'height_mm': mm(section.page_height)}, 'info' if size_ok else 'error')
        margins = {
            "top_mm": mm(section.top_margin),
            "bottom_mm": mm(section.bottom_margin),
            "left_mm": mm(section.left_margin),
            "right_mm": mm(section.right_margin),
        }
        ok = all(21 <= value <= 25 for value in margins.values())
        add(f"第{idx}节页边距 23mm±2mm", "pass" if ok else "review", margins, "error" if not ok else "info")

    paragraphs = nonempty_paragraphs(doc)
    red_headers = [p for _, p in paragraphs if any(
        run.font.color.rgb is not None and str(run.font.color.rgb) == "FF0000"
        for run in p.runs
    )]
    distributed = len(red_headers) == 2 and all(
        p._p.xpath('./w:pPr/w:jc[@w:val="distribute"]') for p in red_headers
    )
    add("红头两行按定稿分散对齐", "pass" if distributed else "review",
        {"red_header_count": len(red_headers), "distributed": distributed},
        "info" if distributed else "error")
    inherited_header_issues = [{'text': p.text, 'indent': indent_issues(p, chars=0),
                                'spacing': spacing_issues(p, line_pt=65)} for p in red_headers]
    inherited_header_issues = [x for x in inherited_header_issues if x['indent'] or x['spacing']]
    add('红头有效缩进为0且实际格式不受网格影响', 'fail' if inherited_header_issues else 'pass',
        inherited_header_issues, 'error' if inherited_header_issues else 'info')
    header_style_ok = len(red_headers) == 2 and all(
        p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
        and pt(p.paragraph_format.line_spacing) == 65
        and all(east_asia_font(r) in ("宋体", "SimSun") and pt(r.font.size) == 45
                and r.bold is True and str(r.font.color.rgb) == "FF0000"
                for r in p.runs if r.text.strip())
        for p in red_headers
    )
    add("红头统一宋体45磅、红色加粗、固定65磅行距",
        "pass" if header_style_ok else "review", header_style_ok,
        "info" if header_style_ok else "error")
    header_text_and_spacing_ok = (
        len(red_headers) == 2
        and red_headers[0].text == "内蒙古准格尔旗"
        and red_headers[1].text == "力量煤业有限公司文件"
        and all(not p._p.xpath('./w:pPr/w:ind') for p in red_headers)
        and all(not p._p.xpath('.//w:rPr/w:spacing') for p in red_headers)
    )
    add("红头原文完整且无手工字距和缩进",
        "pass" if header_text_and_spacing_ok else "review",
        {"red_header_text": [p.text for p in red_headers],
         "manual_tracking": [bool(p._p.xpath('.//w:rPr/w:spacing')) for p in red_headers],
         "paragraph_indents": [bool(p._p.xpath('./w:pPr/w:ind')) for p in red_headers],
         "note": "仍需在WPS/Word逐页核对两行左右齐平及整体居中"},
        "info" if header_text_and_spacing_ok else "error")
    issuing_number = "" if args.blank_document_number else (args.issuing_number or next(
        (
            p.text.strip()
            for _, p in paragraphs
            if re.fullmatch(r".+〔\d{4}〕\d+号", p.text.strip())
            or p.text.strip() == "〔发文字号待补〕"
        ),
        "",
    ))
    number_paragraphs = ([p for _, p in paragraphs if p.text.strip() == issuing_number]
                         if issuing_number else [])
    if args.blank_document_number:
        blank_number = doc.paragraphs[2] if len(doc.paragraphs) > 2 else None
        blank_number_ok = bool(
            blank_number is not None
            and blank_number.text.strip() == ""
            and blank_number.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
            and pt(blank_number.paragraph_format.line_spacing) == args.blank_number_line_pt
        )
        add("文号正文按要求留空且保留设定行距",
            "pass" if blank_number_ok else "review",
            {"paragraph_index": 3, "text": blank_number.text if blank_number is not None else None,
             "line_spacing_pt": pt(blank_number.paragraph_format.line_spacing) if blank_number is not None else None,
             "expected_line_spacing_pt": args.blank_number_line_pt,
             "note": "文件名仍标注发文字号待补；取得正式文号后须重新验收"},
            "info" if blank_number_ok else "error")
    number_font_issues = []
    for p in number_paragraphs:
        for r in p.runs:
            if not r.text.strip():
                continue
            fonts = r._r.rPr.rFonts if r._r.rPr is not None else None
            cjk = fonts.get(qn("w:eastAsia")) if fonts is not None else None
            latin = fonts.get(qn("w:ascii")) if fonts is not None else None
            if any(not c.isascii() for c in r.text) and cjk not in ("宋体", "SimSun"):
                number_font_issues.append({"text": r.text, "chinese_font": cjk})
            if any(c.isdigit() and c.isascii() for c in r.text) and latin != "Times New Roman":
                number_font_issues.append({"text": r.text, "digit_font": latin})
            if pt(r.font.size) != 14 or r.bold is not False:
                number_font_issues.append({"text": r.text, "size_pt": pt(r.font.size), "bold": r.bold})
    if not args.blank_document_number:
        add("文号宋体中文、Times New Roman数字、四号不加粗",
            "pass" if number_paragraphs and not number_font_issues else "review",
            number_font_issues or {"number_paragraphs": len(number_paragraphs)},
            "error" if not number_paragraphs or number_font_issues else "info")
    red_rule_paragraphs = [
        p for p in doc.paragraphs
        if any(b.get(qn("w:color")) == "FF0000" for b in p._p.xpath('./w:pPr/w:pBdr/w:bottom'))
    ]
    if red_rule_paragraphs:
        p = red_rule_paragraphs[0]
        sec = doc.sections[0]
        pf = p.paragraph_format
        width = pt(sec.page_width - sec.left_margin - sec.right_margin -
                   (pf.left_indent or 0) - (pf.right_indent or 0))
        b = p._p.xpath('./w:pPr/w:pBdr/w:bottom')[0]
        ok = abs(width - RED_LINE_WIDTH_PT) <= 2 and b.get(qn("w:sz")) == "16"
        add("红线约178毫米、2磅粗、伸出正文版心", "pass" if ok else "review",
            {"calculated_width_pt": width, "width_mm": round(width * 25.4 / 72, 2), "paragraph": p.text.strip(), "note": "仍须渲染确认最终长度与位置"},
            "info" if ok else "error")
    else:
        add("红线约178毫米、2磅粗、伸出正文版心", "manual", "模板使用图形等非段落边框时需测量渲染效果")

    circulation_prefixes = ("主题：", "抄送：", "抄报：", "发文：", "印发：")
    if args.inspection_material:
        visible_text = "".join(doc._element.body.xpath('.//w:t/text()'))
        furniture = re.findall(r"签\s*发\s*人\s*[:：]|(?:主题|抄报|抄送|发文|印发)\s*[:：]|共印\s*\d+\s*份", visible_text)
        add("迎检文件无签发人及末尾印发栏", "review" if furniture else "pass", furniture,
            "error" if furniture else "info")
        add("迎检文件无遗留印发黑横线/文本框", "manual", "逐页核查原印发区域，含图形线条；不删除红色分隔线")
    elif args.sign_type == "internal":
        signers = [(idx, p) for idx, p in paragraphs if is_signer_paragraph(p)]
        circulation = [(idx, p) for idx, p in paragraphs if p.text.strip().startswith(circulation_prefixes)]
        add("内发文包含签发人栏", "pass" if signers else "review", [idx for idx, _ in signers], "error" if not signers else "info")
        signer_layout_ok = False
        signer_evidence = {}
        if len(signers) == 1:
            signer_paragraph = signers[0][1]
            signer_pos = next(i for i, p in enumerate(doc.paragraphs) if p._p is signer_paragraph._p)
            next_paragraph = doc.paragraphs[signer_pos + 1] if signer_pos + 1 < len(doc.paragraphs) else None
            label_runs = signer_paragraph.runs
            tracking = (label_runs[0]._r.rPr.find(qn("w:spacing"))
                        if label_runs and label_runs[0]._r.rPr is not None else None)
            tracking_twips = int(tracking.get(qn("w:val"))) if tracking is not None else None
            rule_below = bool(next_paragraph and next_paragraph._p.xpath(
                './w:pPr/w:pBdr/w:bottom[@w:color="FF0000"]'))
            signer_evidence = {
                "right_indent_pt": pt(signer_paragraph.paragraph_format.right_indent),
                "expected_right_indent_pt": args.signer_right_indent_pt,
                "label_runs": [r.text for r in label_runs],
                "first_tracking_twips": tracking_twips,
                "expected_tracking_twips": args.signer_label_tracking_twips,
                "wide_rule_on_next_paragraph": rule_below,
                "rule_line_pt": pt(next_paragraph.paragraph_format.line_spacing) if next_paragraph else None,
                "rule_after_pt": pt(next_paragraph.paragraph_format.space_after) if next_paragraph else None,
            }
            signer_layout_ok = (
                signer_paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
                and signer_evidence["right_indent_pt"] is not None
                and abs(signer_evidence["right_indent_pt"] - args.signer_right_indent_pt) <= 0.05
                and len(label_runs) >= 2
                and label_runs[0].text == "签发"
                and label_runs[1].text == "人："
                and tracking_twips == args.signer_label_tracking_twips
                and rule_below
                and signer_evidence["rule_line_pt"] == 2
                and signer_evidence["rule_after_pt"] == 3
            )
        add("内发文签发栏及红线位置按同类样本设置",
            "pass" if signer_layout_ok else "review", signer_evidence,
            "info" if signer_layout_ok else "error")
        add("内发文包含主题、抄送、发文和共印印发栏", "pass" if len(circulation) == 3 and any("共印" in p.text for _, p in circulation) else "review",
            {"paragraphs": [{"index": idx, "text": p.text} for idx, p in circulation]},
            "error" if len(circulation) != 3 or not any("共印" in p.text for _, p in circulation) else "info")
        if circulation:
            first, last = circulation[0][1], circulation[-1][1]
            has_top = bool(first._p.xpath('./w:pPr/w:pBdr/w:top[@w:val="single"]'))
            has_bottom = bool(last._p.xpath('./w:pPr/w:pBdr/w:bottom[@w:val="single"]'))
            add("内发印发栏上下为黑色实线", "pass" if has_top and has_bottom else "review",
                {"top": has_top, "bottom": has_bottom}, "error" if not (has_top and has_bottom) else "info")
    footer_issues = []
    for i, section in enumerate(doc.sections, 1):
        footer = section.footer
        page_runs = [r for p in footer.paragraphs for r in p.runs if "PAGE" in r._r.xml]
        if not page_runs:
            footer_issues.append({"section": i, "issue": "missing_PAGE"})
        for r in page_runs:
            if r.font.name != "Times New Roman" or pt(r.font.size) != 9:
                footer_issues.append({"section": i, "font": r.font.name, "size_pt": pt(r.font.size)})
    add("页码Times New Roman、9磅、PAGE域", "pass" if not footer_issues else "review", footer_issues, "error" if footer_issues else "info")
    if not paragraphs:
        add("存在正文段落", "fail", "未找到非空段落", "error")
    else:
        add("存在正文段落", "pass", f"共 {len(paragraphs)} 个非空段落")

    title_hits = []
    if args.title:
        title_hits = [(idx, p) for idx, p in paragraphs if canonical_title(p.text) == canonical_title(args.title)]
        add("正文标题与传入标题一致", "pass" if title_hits else "review", [idx for idx, _ in title_hits], "error" if not title_hits else "info")
        expected_name = official_filename(args.title, issuing_number)
        name_ok = args.docx.name == expected_name
        add("文件名为发文字号：正文标题.docx", "pass" if name_ok else "review", {"filename": args.docx.name, "expected": expected_name}, "error" if not name_ok else "info")
        add("多行标题上长下短且不拆词", "manual",
            {"title_lines": [p.text.splitlines() for _, p in title_hits], "note": "需渲染核对各行实际宽度及语义边界，不能只比较字数"})
        add("标题汉字数字间约四分之一字宽间隔", "manual",
            "检查渲染后间距；autoSpaceDN/DE仅是复现方式，不保证实际字形间隔")
        title_issues = []
        for idx, paragraph in title_hits:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                effective_issues = run_format_issues(run, font='黑体', size=22, bold=False)
                if effective_issues:
                    title_issues.append({'paragraph': idx, 'text': run.text[:30], **effective_issues})
                font_name = east_asia_font(run) or run.font.name
                size = pt(run.font.size)
                if font_name and "黑体" not in font_name and "SimHei" not in font_name:
                    title_issues.append({"paragraph": idx, "font": font_name, "text": run.text[:20]})
                if size is not None and abs(size - 22) > 0.5:
                    title_issues.append({"paragraph": idx, "size_pt": size, "text": run.text[:20]})
                if run.bold is True:
                    title_issues.append({"paragraph": idx, "bold": True, "text": run.text[:20]})
        add("标题黑体二号且不加粗", "pass" if title_hits and not title_issues else "review", title_issues[:20] if title_issues else [idx for idx, _ in title_hits], "error" if title_issues else "info")
        title_spacing_issues = []
        for idx, paragraph in title_hits:
            before = pt(paragraph.paragraph_format.space_before)
            after = pt(paragraph.paragraph_format.space_after)
            if before not in (None, 0.0):
                title_spacing_issues.append({"paragraph": idx, "space_before_pt": before})
            if after not in (None, 0.0):
                title_spacing_issues.append({"paragraph": idx, "space_after_pt": after})
        add("标题段前段后均为0", "pass" if title_hits and not title_spacing_issues else "review", title_spacing_issues[:20] if title_spacing_issues else [idx for idx, _ in title_hits], "error" if title_spacing_issues else "info")
        title_line_spacing_issues = []
        for idx, paragraph in title_hits:
            spacing = paragraph._p.pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}spacing") if paragraph._p.pPr is not None else None
            if spacing is not None:
                for attribute in ("before", "after", "beforeLines", "afterLines", "beforeAutospacing", "afterAutospacing"):
                    value = spacing.get(f"{{http://schemas.openxmlformats.org/wordprocessingml/2006/main}}{attribute}")
                    if value is not None:
                        title_line_spacing_issues.append({"paragraph": idx, attribute: value})
        add("标题段前段后继承模板0行显示", "pass" if title_hits and not title_line_spacing_issues else "review", title_line_spacing_issues[:20] if title_line_spacing_issues else [idx for idx, _ in title_hits], "error" if title_line_spacing_issues else "info")
        title_style_issues = []
        for idx, paragraph in title_hits:
            if paragraph.style.name != "Normal":
                title_style_issues.append({"paragraph": idx, "style": paragraph.style.name, "expected": "Normal"})
        add("标题使用正文文本Normal样式", "pass" if title_hits and not title_style_issues else "review", title_style_issues[:20] if title_style_issues else [idx for idx, _ in title_hits], "error" if title_style_issues else "info")
        inherited_title_issues = []
        for idx, paragraph in title_hits:
            problem = spacing_issues(paragraph, single=True)
            indent = indent_issues(paragraph, chars=0)
            if problem or indent:
                inherited_title_issues.append({'paragraph': idx, 'spacing': problem, 'indent': indent})
        add('标题继承后的段前段后为0、单倍且无遗留缩进', 'fail' if inherited_title_issues else 'pass',
            inherited_title_issues, 'error' if inherited_title_issues else 'info')

    recipient_paragraph = None if args.no_recipient else next(
        (
            p
            for idx, p in paragraphs
            if not (args.title and canonical_title(p.text) == canonical_title(args.title))
            and p.text.strip() != issuing_number
            and not is_signer_paragraph(p)
            and not has_paragraph_border(p)
            and p not in red_headers
        ),
        None,
    )
    if args.no_recipient:
        add('原稿无主送分支', 'pass', '已显式传入--no-recipient；首段按正文检查')
    elif recipient_paragraph is None or not recipient_paragraph.text.strip().endswith(('：', ':')):
        add('主送单位可明确识别', 'fail', '首段不是明确的主送段；核对原稿，确无主送时生成与校验均传--no-recipient', 'error')
    body_candidates = [
        (idx, p)
        for idx, p in paragraphs
        if not (args.title and canonical_title(p.text) == canonical_title(args.title))
        and p.text.strip() != issuing_number
        and not is_signer_paragraph(p)
        and not p.text.strip().startswith(circulation_prefixes)
    ]
    font_issues = []
    recipient_font_issues = []
    heading_font_issues = []
    paragraph_issues = []
    semicolon_hits = []
    signatures = {p._p for _, p in paragraphs if p.text.strip() in (args.expected_company, args.expected_mine, '大饭铺煤矿')
                  or re.fullmatch(r'\d{4}年\d{1,2}月\d{1,2}日', p.text.strip())}
    for idx, paragraph in body_candidates:
        # Red-head paragraphs are intentionally not body text.
        if has_paragraph_border(paragraph) or paragraph in red_headers:
            continue
        if paragraph.text.rstrip().endswith("；"):
            semicolon_hits.append(idx)
        pf = paragraph.paragraph_format
        spacing_problem = spacing_issues(paragraph)
        if on(effective_ppr(paragraph, 'snapToGrid').get('val')):
            spacing_problem['snapToGrid'] = 'enabled'
        indent_problem = {} if paragraph._p in signatures else indent_issues(paragraph, chars=0 if paragraph is recipient_paragraph else 2)
        if spacing_problem or indent_problem:
            paragraph_issues.append({'paragraph': idx, 'spacing': spacing_problem, 'indent': indent_problem})
        allow_adjustment = args.allow_table_spacing_adjustment and bool(doc.tables)
        line_pt = pt(pf.line_spacing) if pf.line_spacing_rule == WD_LINE_SPACING.EXACTLY else None
        if line_pt is None or not (25 <= line_pt <= 28 if allow_adjustment else abs(line_pt - 28) < 0.05):
            paragraph_issues.append({"paragraph": idx, "line_spacing_pt": line_pt, "expected": "25—28磅有表格例外" if allow_adjustment else "固定28磅"})
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            is_heading = independent_heading(paragraph.text)
            expected_bold = True if paragraph is recipient_paragraph or is_heading else False if paragraph._p in signatures else args.body_bold
            target = recipient_font_issues if paragraph is recipient_paragraph else heading_font_issues if is_heading else font_issues
            issues = run_format_issues(run, font='仿宋', size=14, bold=expected_bold)
            if issues:
                target.append({'paragraph': idx, 'text': run.text[:30], **issues})
        marker = paragraph._p.xpath('./w:pPr/w:rPr')
        if marker:
            bold = marker[0].find(qn('w:b'))
            if bold is not None:
                expected = paragraph is recipient_paragraph or independent_heading(paragraph.text) or (args.body_bold and paragraph._p not in signatures)
                if on(bold.get(qn('w:val'), '1')) != expected:
                    paragraph_issues.append({'paragraph': idx, 'paragraph_marker_bold': 'conflicts with paragraph role'})

    body_label = "普通正文仿宋四号且按当次要求加粗" if args.body_bold else "普通正文仿宋四号且不加粗"
    add(body_label, "pass" if not font_issues else "review", font_issues[:20], "error" if font_issues else "info")
    add("主送单位仿宋四号且加粗", "pass" if not recipient_font_issues else "review", recipient_font_issues[:20], "error" if recipient_font_issues else "info")
    add('一级独立小标题仿宋四号且加粗', 'fail' if heading_font_issues else 'pass', heading_font_issues,
        'error' if heading_font_issues else 'info')
    add("正文段前段后为0、默认固定28磅（有表格局部压缩须显式声明）", "pass" if not paragraph_issues else "review", paragraph_issues[:20], "error" if paragraph_issues else "info")
    add("正文条目不以分号结尾", "pass" if not semicolon_hits else "review", semicolon_hits, "error" if semicolon_hits else "info")

    full_text = "\n".join(p.text for _, p in paragraphs)
    book_hits = [{'paragraph': idx, 'text': p.text[:120]} for idx, p in paragraphs if re.search(r'》\s*、\s*《', p.text)]
    add('相邻书名号之间无顿号', 'fail' if book_hits else 'pass', book_hits, 'error' if book_hits else 'info')
    hierarchy_issues, punctuation_issues, previous_level, in_attachments = [], [], None, False
    for idx, p in body_candidates:
        text = p.text.strip()
        if p in red_headers or has_paragraph_border(p) or p is recipient_paragraph or p._p in signatures:
            continue
        if attachment_start(text):
            in_attachments = True
            continue
        if closing(text):
            in_attachments = False
            continue
        if in_attachments and re.match(r'^\d+[.．]', text):
            continue
        in_attachments = False
        level = numbered_level(text)
        if level is not None:
            if (previous_level is None and level > 1) or (previous_level is not None and level > previous_level + 1):
                hierarchy_issues.append({'paragraph': idx, 'text': text[:80], 'previous_level': previous_level, 'level': level})
            previous_level = level
            if not independent_heading(text) and not text.endswith(('。', '：', ':')):
                punctuation_issues.append({'paragraph': idx, 'text': text[-100:], 'expected': '完整条目以。结束；冒号引出内容和独立小标题除外'})
    add('显式序号层级无跳级', 'fail' if hierarchy_issues else 'pass', hierarchy_issues, 'error' if hierarchy_issues else 'info')
    add('完整序号条目用句号结束', 'fail' if punctuation_issues else 'pass', punctuation_issues, 'error' if punctuation_issues else 'info')
    if any(p._p.xpath('./w:pPr/w:numPr[not(w:numId/@w:val="0")]') for _, p in paragraphs):
        add('自动编号的实际显示层级与序号', 'manual', 'XML自动编号需结合numbering.xml与实际页面核对，未按纯文本规则假定通过')
    closing_issues = [{'paragraph': idx, 'text': p.text} for idx, p in paragraphs if closing(p.text) and p.text.strip()[-1] in '。.;；！!？?']
    add('已识别结束语不带结束标点', 'fail' if closing_issues else 'pass', closing_issues, 'error' if closing_issues else 'info')
    separator_issues, protected_separators = [], []
    all_p = list(doc.paragraphs)
    for i, p in enumerate(all_p):
        if not attachment_start(p.text) and not closing(p.text) and p.text.strip() != args.expected_company:
            continue
        j = i - 1
        blanks = 0
        while j >= 0 and safe_blank(all_p[j]):
            blanks += 1
            j -= 1
        if j >= 0 and not all_p[j].text.strip():
            protected_separators.append({'paragraph': j + 1, 'before': p.text[:80]})
        applicable = p.text.strip() == args.expected_company or (attachment_start(p.text) and j >= 0 and closing(all_p[j].text))
        if closing(p.text) and j >= 0 and re.match(r'^\d+[.．]', all_p[j].text.strip()):
            applicable = any(attachment_start(x.text) for x in all_p[:j + 1])
        if applicable and blanks != 1:
            separator_issues.append({'paragraph': i + 1, 'text': p.text[:80], 'blank_count': blanks, 'expected': 1})
    add('已识别附件/结束语及落款前恰好空一行', 'fail' if separator_issues else 'pass', separator_issues, 'error' if separator_issues else 'info')
    if protected_separators:
        add('带标记或对象空段的实际分隔间距', 'manual', protected_separators)
    companies = [p for _, p in paragraphs if p.text.strip() == args.expected_company]
    dates = [p for _, p in paragraphs if re.fullmatch(r"\d{4}年\d{1,2}月\d{1,2}日", p.text.strip())]
    date_issues = []
    for p in dates:
        match = re.fullmatch(r'(\d{4})年(\d{1,2})月(\d{1,2})日', p.text.strip())
        try:
            calendar_date(int(match[1]), int(match[2]), int(match[3]))
        except ValueError:
            date_issues.append({'text': p.text, 'issue': 'invalid_calendar_date'})
        if match[2].startswith('0') or match[3].startswith('0'):
            date_issues.append({'text': p.text, 'issue': 'leading_zero'})
    add('落款日期为有效阿拉伯日期且无前导0', 'fail' if not dates or date_issues else 'pass',
        date_issues or {'date_count': len(dates)}, 'error' if not dates or date_issues else 'info')
    mine = args.expected_mine or ('大饭铺煤矿' if '大饭铺煤矿' in full_text else None)
    separate_signature = bool(companies) and (mine is None or any(p.text.strip() == mine for _, p in paragraphs))
    add('公司与矿名在独立段落分行', 'pass' if separate_signature else 'fail',
        {'company_paragraph_count': len(companies), 'mine': mine}, 'info' if separate_signature else 'error')
    signature_ok = bool(companies and dates) and all(
        p._p.xpath('./w:pPr/w:jc[@w:val="center"]') and
        p.paragraph_format.left_indent == companies[-1].paragraph_format.left_indent and
        p.paragraph_format.right_indent == companies[-1].paragraph_format.right_indent
        for p in (companies[-1], dates[-1])
    )
    add("公司落款与日期在同一对齐区域居中", "pass" if signature_ok else "review",
        {"company_found": bool(companies), "date_found": bool(dates), "note": "结构校验不能替代实际字形居中的渲染检查"},
        "info" if signature_ok else "error")
    expected_values = [("公司名", args.expected_company)]
    if args.expected_mine:
        expected_values.append(("矿名", args.expected_mine))
    for label, value in expected_values:
        add(f"包含{label}", "pass" if value in full_text else "review", value, "error" if value not in full_text else "info")

    table_count, border_issues = table_border_issues(args.docx)
    add(
        "所有表格边框均为single实线",
        "pass" if table_count == 0 or not border_issues else "review",
        {"table_count": table_count, "issues": border_issues[:30]},
        "error" if border_issues else "info",
    )

    if not args.title:
        add("标题黑体二号且不加粗", "manual", "请传入 --title 后自动检查标题段落")
    add("红头颜色、红线、右下角落款、页码和上长下短", "manual", "需渲染页面图片或用目标办公软件目测确认")
    add("公司落款靠近版心右边界，矿名和日期以公司为基准居中", "manual", "核对公司名称实际字宽；不得仅凭固定220磅左缩进或与旧版相比平移5毫米验收")
    add('序号语义、附件范围及未知结束语', 'manual', '机械规则已独立检查；完整词语、条目是否完整与附件业务范围仍需逐页确认')

    failed = sum(c['status'] == 'fail' for c in checks)
    pending = sum(c['status'] in ('review', 'manual') for c in checks)
    report = {"file": str(args.docx), "checks": checks, "summary": {"total": len(checks),
              "passed": sum(c['status'] == 'pass' for c in checks), "review": pending, "pending_manual": pending,
              "failed": failed, "automatic_ok": failed == 0, "formal_ready": failed == 0 and pending == 0}}
    if args.json:
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        print(f"文件：{args.docx}")
        for check in checks:
            print(f"[{check['status']}] {check['rule']}：{check['evidence']}")
        print(f"汇总：{report['summary']}")
    return 1 if failed else 2 if pending and args.require_manual_review else 0


if __name__ == "__main__":
    raise SystemExit(main())
