#!/usr/bin/env python3
"""Apply the confirmed real-execution internal layout to a copy of an OOXML file."""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re
from zipfile import ZipFile, is_zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.shared import Pt

from apply_red_header_docx import (set_spacing, set_exact_spacing, set_run, reset_indent,
                                  trim_recipient_leading_whitespace, official_filename,
                                  trim_post_signature_blank_pages, set_paragraph_border)
from red_header_common import (effective_bold, text_runs, safe_blank, independent_heading,
                               effective_ppr, on, indent_issues, run_format_issues)

CLOSING_PUNCTUATION = '，。；：、！？）》】〕〉」』”’％‰℃～~,.;:!?)]}'


def set_flag(paragraph, tag, enabled):
    ppr = paragraph._p.get_or_add_pPr()
    node = ppr.find(qn('w:' + tag))
    if node is None:
        node = OxmlElement('w:' + tag)
        ppr.append(node)
    node.set(qn('w:val'), '1' if enabled else '0')


def prevent_leading_punctuation(paragraph):
    set_flag(paragraph, 'kinsoku', True)
    set_flag(paragraph, 'overflowPunct', False)
    set_flag(paragraph, 'topLinePunct', False)


def replace_character(paragraph, offset, replacement):
    for node in paragraph._p.iter(qn('w:t')):
        value = node.text or ''
        if offset < len(value):
            node.text = value[:offset] + replacement + value[offset + 1:]
            return
        offset -= len(value)
    raise ValueError('无法安全定位要修正的序号或标点。')


def subheading(text):
    return bool(re.fullmatch(r'[（(][一二三四五六七八九十百]+[）)][^，。；：！？,.;:!?]{1,45}[:：]?', text))


def normalize_execution_paragraph(paragraph):
    trim_recipient_leading_whitespace(paragraph)
    match = re.match(r'^(\d+)、', paragraph.text)
    if match:
        replace_character(paragraph, len(match[1]), '.')
    if subheading(paragraph.text.strip()) and paragraph.text.endswith(('：', ':')):
        replace_character(paragraph, len(paragraph.text) - 1, '')


def sync_marker(paragraph):
    first = next((r for r in text_runs(paragraph) if r.text.strip()), None)
    if first is None or first._r.rPr is None:
        return
    ppr = paragraph._p.get_or_add_pPr()
    for node in list(ppr.findall(qn('w:rPr'))):
        ppr.remove(node)
    ppr.append(deepcopy(first._r.rPr))


def table_shape(table):
    return {'texts': table._tbl.xpath('.//w:t/text()'),
            'grid': [dict(n.attrib) for n in table._tbl.xpath('./w:tblGrid/w:gridCol')],
            'rows': len(table.rows),
            'cells': [len(r._tr.tc_lst) for r in table.rows],
            'merges': [(n.tag, dict(n.attrib)) for n in table._tbl.xpath('.//w:vMerge | .//w:gridSpan')]}


def revise_tables(doc, keep_together_tables=()):
    for ti, table in enumerate(doc.tables, 1):
        before = table_shape(table)
        for i, row in enumerate(table.rows):
            if i == 0 and not row._tr.xpath('./w:trPr/w:tblHeader'):
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for tc in row._tr.tc_lst:
                cell = _Cell(tc, table)
                if i == 0:
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for p in cell.paragraphs:
                    set_exact_spacing(p, 22)
                    reset_indent(p, chars=0)
                    prevent_leading_punctuation(p)
                    p.paragraph_format.keep_with_next = i == 0 or (ti in keep_together_tables and i < len(table.rows) - 1)
                    p.paragraph_format.keep_together = False
                    p.paragraph_format.page_break_before = False
                    if i == 0:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in text_runs(p):
                        set_run(r, font='仿宋', size=12, bold=True if i == 0 else effective_bold(r))
                    sync_marker(p)
        for edge in table._tbl.xpath('./w:tblPr/w:tblBorders/* | .//w:tcPr/w:tcBorders/*'):
            if edge.get(qn('w:val')) not in ('nil', 'none'):
                edge.set(qn('w:val'), 'single')
        previous = table._tbl.getprevious()
        while previous is not None and previous.tag == qn('w:p'):
            p = Paragraph(previous, doc._body)
            if not safe_blank(p):
                break
            previous = previous.getprevious()
            p._p.getparent().remove(p._p)
        if previous is not None and previous.tag == qn('w:p'):
            caption = Paragraph(previous, doc._body)
            if re.match(r'^表\s*\d+', caption.text.strip()):
                caption.paragraph_format.keep_with_next = True
                caption.paragraph_format.keep_together = True
                caption.paragraph_format.page_break_before = False
        if table_shape(table) != before:
            raise ValueError('表格文字、列网格或合并结构发生了非授权改变。')


def distribution_values(template, title):
    copy_to = str(template.get('copy_to', '')).strip()
    issue_unit = str(template.get('issue_unit', '')).strip()
    count = template.get('print_count')
    if not copy_to or not issue_unit or isinstance(count, bool) or not isinstance(count, int) or count <= 0:
        raise ValueError('固定印发模板须提供copy_to、issue_unit及正整数print_count；不自动猜测。')
    return [f'主题：{title}', f'抄送：{copy_to}', f'发文：{issue_unit}\t共印{count}份']


def revise_distribution(doc, title, template):
    expected = distribution_values(template, title)
    block = [p for p in doc.paragraphs if p.text.strip().startswith(('主题：', '抄送：', '发文：'))]
    if len(block) != 3:
        raise ValueError('固定印发模板只用于已明确识别的三行印发栏，不删除未知内容。')
    for i, (p, text) in enumerate(zip(block, expected)):
        p.clear()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_exact_spacing(p, 17)
        reset_indent(p, chars=0)
        p.paragraph_format.keep_with_next = i < 2
        p.paragraph_format.keep_together = True
        set_run(p.add_run(text), font='仿宋', size=12, bold=False)
        sync_marker(p)
    usable = doc.sections[-1].page_width - doc.sections[-1].left_margin - doc.sections[-1].right_margin
    block[-1].paragraph_format.tab_stops.clear_all()
    block[-1].paragraph_format.tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
    set_paragraph_border(block[0], 'top', size='8', space='2')
    set_paragraph_border(block[-1], 'bottom', size='8', space='2')


def revise(input_path, output_path, *, title, distribution_template=None, keep_together_tables=()):
    if not is_zipfile(input_path):
        raise ValueError('输入不是可编辑的Word OOXML包；网页文件不能只改扩展名交付。')
    with ZipFile(input_path) as z:
        if 'word/document.xml' not in z.namelist():
            raise ValueError('压缩包中不存在Word正文，不能当DOCX处理。')
    doc = Document(input_path)
    if any(i < 1 or i > len(doc.tables) for i in keep_together_tables):
        raise ValueError('整表同页的序号超出原稿表格范围。')
    settings = doc.settings.element
    for language in ('zh-CN', 'en-US'):
        custom = next((n for n in settings.findall(qn('w:noLineBreaksBefore')) if n.get(qn('w:lang')) == language), None)
        if custom is None:
            custom = OxmlElement('w:noLineBreaksBefore')
            custom.set(qn('w:lang'), language)
            settings.append(custom)
        custom.set(qn('w:val'), ''.join(dict.fromkeys(custom.get(qn('w:val'), '') + CLOSING_PUNCTUATION)))
    numbers = list(dict.fromkeys(p.text.strip() for p in doc.paragraphs
                                if re.fullmatch(r'蒙力煤字〔\d{4}〕\d+号', p.text.strip())))
    if len(numbers) != 1:
        raise ValueError('须从原稿确认唯一的内发文号，不自动编造。')
    output_path = output_path.with_name(official_filename(title, numbers[0]))
    if input_path.resolve() == output_path.resolve():
        raise ValueError('禁止覆盖源文件。')
    paragraphs = list(doc.paragraphs)
    start = next((i for i, p in enumerate(paragraphs) if p.text.strip() == title), None)
    end = next((i for i, p in enumerate(paragraphs) if i > (start if start is not None else -1)
                and p.text.strip() == '内蒙古准格尔旗力量煤业有限公司'), None)
    if start is None or end is None:
        raise ValueError('正文标题或落款边界不明确，须先核对。')
    changes = []
    for p in paragraphs[start + 1:end]:
        if not p.text.strip():
            continue
        old = p.text
        normalize_execution_paragraph(p)
        set_spacing(p, indent_pt=28)
        # A first-line indent must not coexist with a hanging-indent override.
        ind = p._p.get_or_add_pPr().find(qn('w:ind'))
        for key in ('hanging', 'hangingChars'):
            ind.attrib.pop(qn('w:' + key), None)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        prevent_leading_punctuation(p)
        p.paragraph_format.keep_with_next = independent_heading(p.text) or subheading(p.text.strip())
        p.paragraph_format.page_break_before = False
        for r in text_runs(p):
            set_run(r, font='仿宋', size=14, bold=effective_bold(r))
        sync_marker(p)
        if old != p.text:
            changes.append({'before': old, 'after': p.text})
    revise_tables(doc, keep_together_tables)
    if distribution_template is not None:
        revise_distribution(doc, title, distribution_template)
    trim_post_signature_blank_pages(doc)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path, changes


def execution_issues(doc, title, template=None, keep_together_tables=()):
    issues = []
    paragraphs = list(doc.paragraphs)
    start = next((i for i, p in enumerate(paragraphs) if p.text.strip() == title), None)
    end = next((i for i, p in enumerate(paragraphs) if i > (start if start is not None else -1)
                and p.text.strip() == '内蒙古准格尔旗力量煤业有限公司'), None)
    if start is None or end is None:
        return [{'issue': '正文标题或落款边界不明确'}]
    for p in paragraphs[start + 1:end]:
        if not p.text.strip():
            continue
        if re.match(r'^\s*\d+、', p.text):
            issues.append({'issue': '阿拉伯序号仍使用顿号', 'text': p.text[:50]})
        if subheading(p.text.strip()) and p.text.rstrip().endswith(('：', ':')):
            issues.append({'issue': '独立括号小标题仍带冒号', 'text': p.text})
        if p.text[:1].isspace() or indent_issues(p, chars=2):
            issues.append({'issue': '正文或小标题未按两字符首行缩进', 'text': p.text[:50]})
        if p.text.lstrip()[:1] in CLOSING_PUNCTUATION:
            issues.append({'issue': '闭合标点位于段首', 'text': p.text[:50]})
        if not on(effective_ppr(p, 'kinsoku').get('val')):
            issues.append({'issue': '未启用中文行首尾禁则', 'text': p.text[:50]})
    for i, table in enumerate(doc.tables, 1):
        for ri, row in enumerate(table.rows):
            for tc in row._tr.tc_lst:
                cell = _Cell(tc, table)
                for p in cell.paragraphs:
                    spacing = effective_ppr(p, 'spacing')
                    if spacing.get('lineRule') != 'exact' or spacing.get('line') != '440':
                        issues.append({'table': i, 'row': ri + 1, 'issue': '表内行距不是固定22磅'})
                    if ri == 0 and (effective_ppr(p, 'jc').get('val') != 'center' or cell.vertical_alignment != WD_CELL_VERTICAL_ALIGNMENT.CENTER):
                        issues.append({'table': i, 'issue': '表头未水平及垂直居中'})
                    if ri > 0 and p.paragraph_format.keep_with_next and i not in keep_together_tables:
                        issues.append({'table': i, 'row': ri + 1, 'issue': '表内存在连续与下段同页链'})
                    for r in text_runs(p):
                        if r.text.strip() and run_format_issues(r, font='仿宋', size=12, bold=True if ri == 0 else None):
                            issues.append({'table': i, 'row': ri + 1, 'issue': '表内字体/字号或表头加粗不符', 'text': r.text[:30]})
        previous = table._tbl.getprevious()
        if previous is not None and previous.tag == qn('w:p'):
            p = Paragraph(previous, doc._body)
            if re.match(r'^表\s*\d+', p.text.strip()) and not p.paragraph_format.keep_with_next:
                issues.append({'table': i, 'issue': '表题未与表头相随'})
    if template is not None:
        expected = distribution_values(template, title)
        actual = [p.text for p in doc.paragraphs if p.text.strip().startswith(('主题：', '抄送：', '发文：'))]
        if actual != expected:
            issues.append({'issue': '印发栏与用户已确认的固定模板不一致', 'actual': actual, 'expected': expected})
    return issues


def main():
    parser = argparse.ArgumentParser(description='修正真实执行内发文表格、序号、行首标点、缩进和固定印发栏')
    parser.add_argument('input', type=Path)
    parser.add_argument('output', type=Path)
    parser.add_argument('--title', required=True)
    parser.add_argument('--distribution-template', type=Path, help='用户已确认的固定copy_to、issue_unit、print_count JSON')
    parser.add_argument('--change-report', type=Path)
    parser.add_argument('--keep-table-together', type=int, action='append', default=[], help='经独立渲染确认可容纳一页的小表序号；长表不得使用，可重复传入')
    args = parser.parse_args()
    template = json.loads(args.distribution_template.read_text()) if args.distribution_template else None
    output, changes = revise(args.input, args.output, title=args.title, distribution_template=template,
                              keep_together_tables=args.keep_table_together)
    if args.change_report:
        args.change_report.write_text(json.dumps({'output': str(output), 'text_changes': changes}, ensure_ascii=False, indent=2) + '\n')
    print(output)


if __name__ == '__main__':
    main()
