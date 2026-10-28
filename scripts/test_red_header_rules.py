#!/usr/bin/env python3
"""Synthetic regression checks; never load business documents or employee data."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from apply_red_header_docx import apply, set_paragraph_border
from extract_review_comments import extract
from validate_outputs import visible_files
from red_header_common import (bold_snapshot, compact_snapshot, effective_bold, text_runs,
                               recipient_layout_issues, rendered_page_content)
from revise_internal_execution_docx import revise as revise_execution, execution_issues, table_shape


TITLE = "关于测试事项的通知"
NUMBER = "蒙力煤字〔2026〕1号"
COMPANY = "内蒙古准格尔旗力量煤业有限公司"


class HeaderRulesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def source(self, decorated=False):
        doc = Document()
        if decorated:
            for text in ("内蒙古准格尔旗", "力量煤业有限公司文件", NUMBER, "签 发 人："):
                doc.add_paragraph(text)
        doc.add_paragraph(TITLE)
        doc.add_paragraph("公司各部门：")
        doc.add_paragraph("一、按已确认要求办理。")
        doc.add_paragraph(COMPANY)
        date = doc.add_paragraph("2026 年 09 月 01 日")
        date.paragraph_format.left_indent = Pt(20)
        if decorated:
            first = doc.add_paragraph("主题：" + TITLE)
            set_paragraph_border(first, "top")
            doc.add_paragraph("抄报：测试范围")
            last = doc.add_paragraph("发文：测试单位\t共印1份")
            set_paragraph_border(last, "bottom")
        path = self.root / "source.docx"
        doc.save(path)
        return path

    def generate(self, decorated=False, **kwargs):
        source = self.source(decorated)
        original = source.read_bytes()
        issuing_number = kwargs.pop("issuing_number", NUMBER)
        output = apply(source, self.root / "output" / "out.docx", title=TITLE,
                       issuing_number=issuing_number, **kwargs)
        self.assertEqual(source.read_bytes(), original)
        return output, Document(output)

    def test_inspection_removes_existing_furniture(self):
        output, doc = self.generate(decorated=True, sign_type="internal", inspection_material=True)
        texts = [p.text for p in doc.paragraphs]
        self.assertEqual(texts.count("内蒙古准格尔旗"), 1)
        self.assertEqual(texts.count(NUMBER), 1)
        self.assertFalse(any(any(word in t for word in ("签 发 人", "主题：", "抄报：", "共印")) for t in texts))
        self.assertIn("一、按已确认要求办理。", texts)
        self.assertFalse(doc._element.body.xpath('.//w:pBdr/*[@w:color="000000"]'))
        self.check_alignment(doc)
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_red_header_docx.py")),
                                 str(output), "--title", TITLE, "--sign-type", "internal",
                                 "--inspection-material", "--json"], check=True, capture_output=True, text=True)
        checks = json.loads(result.stdout)["checks"]
        self.assertFalse([c for c in checks if c["severity"] == "error"], checks)

    def check_alignment(self, doc):
        for p in doc.paragraphs[:2]:
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.DISTRIBUTE)
            self.assertEqual(p.paragraph_format.line_spacing, Pt(65))
            self.assertFalse(p._p.xpath('./w:pPr/w:ind'))
            for r in p.runs:
                self.assertEqual(r.font.name, "宋体")
                self.assertEqual(r.font.size, Pt(45))
                self.assertIsNone(r._r.rPr.find(qn("w:spacing")))
        self.assertEqual(doc.paragraphs[0].text, "内蒙古准格尔旗")
        self.assertEqual(doc.paragraphs[1].text, "力量煤业有限公司文件")
        company = next(p for p in doc.paragraphs if p.text == COMPANY)
        date = next(p for p in doc.paragraphs if p.text == "2026年9月1日")
        self.assertEqual(company.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(date.alignment, WD_ALIGN_PARAGRAPH.CENTER)
        self.assertEqual(company.paragraph_format.left_indent, date.paragraph_format.left_indent)
        self.assertEqual(company.paragraph_format.right_indent, date.paragraph_format.right_indent)

    def test_execution_internal_retains_furniture(self):
        _, doc = self.generate(sign_type="internal", copy_to="测试范围", issue_unit="测试单位", print_count=1)
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("签发人：", text)
        self.assertIn("抄送：测试范围", text)
        self.assertIn("共印1份", text)
        self.check_alignment(doc)

    def test_missing_print_count_stays_unconfirmed(self):
        _, doc = self.generate(sign_type="internal", copy_to="测试范围", issue_unit="测试单位")
        self.assertIn("共印〔印数待确认〕份", "\n".join(p.text for p in doc.paragraphs))

    def test_signer_sample_override_generates_and_validates(self):
        output, doc = self.generate(sign_type="internal", copy_to="测试范围", issue_unit="测试单位",
                                    signer_label_tracking_twips=100, signer_right_indent_pt=80)
        signer = next(p for p in doc.paragraphs if p.text.startswith("签发人："))
        self.assertEqual(signer.paragraph_format.right_indent, Pt(80))
        spacing = signer.runs[0]._r.rPr.find(qn("w:spacing"))
        self.assertEqual(int(spacing.get(qn("w:val"))), 100)
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_red_header_docx.py")),
                                 str(output), "--title", TITLE, "--sign-type", "internal",
                                 "--signer-label-tracking-twips", "100", "--signer-right-indent-pt", "80", "--json"],
                                check=True, capture_output=True, text=True)
        checks = json.loads(result.stdout)["checks"]
        self.assertTrue(any(c["rule"] == "内发文签发栏及红线位置按同类样本设置"
                            and c["status"] == "pass" for c in checks))

    def test_inspection_needs_no_distribution_fields(self):
        self.generate(sign_type="internal", inspection_material=True)

    def test_execution_internal_still_requires_distribution_fields(self):
        with self.assertRaisesRegex(ValueError, "必须提供"):
            self.generate(sign_type="internal")

    def test_validator_rejects_prior_center_plus_tracking(self):
        output, doc = self.generate(inspection_material=True)
        doc.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.paragraphs[1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "450")
        doc.paragraphs[0].runs[0]._r.rPr.append(spacing)
        doc.save(output)
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_red_header_docx.py")),
                                 str(output), "--title", TITLE, "--inspection-material", "--json"],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        self.assertGreater(json.loads(result.stdout)['summary']['failed'], 0)
        checks = json.loads(result.stdout)["checks"]
        self.assertTrue(any(c["rule"] == "红头两行按定稿分散对齐" and c["severity"] == "error" for c in checks))
        self.assertTrue(any(c["rule"] == "红头原文完整且无手工字距和缩进"
                            and c["severity"] == "error" for c in checks))

    def test_validator_rejects_manual_header_tracking(self):
        output, doc = self.generate(sign_type="internal", inspection_material=True)
        first_run = doc.paragraphs[0].runs[0]
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:val"), "450")
        first_run._r.rPr.append(spacing)
        doc.save(output)
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_red_header_docx.py")),
                                 str(output), "--title", TITLE, "--sign-type", "internal", "--inspection-material", "--json"],
                                 capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        checks = json.loads(result.stdout)["checks"]
        self.assertTrue(any(c["rule"] == "红头原文完整且无手工字距和缩进"
                            and c["severity"] == "error" for c in checks))

    def test_external_uses_approved_distributed_header(self):
        output, doc = self.generate(sign_type="external")
        self.check_alignment(doc)
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_red_header_docx.py")),
                                 str(output), "--title", TITLE, "--sign-type", "external", "--json"],
                                check=True, capture_output=True, text=True)
        checks = json.loads(result.stdout)["checks"]
        self.assertTrue(any(c["rule"] == "红头原文完整且无手工字距和缩进"
                            and c["status"] == "pass" for c in checks))

    def test_requested_blank_number_line_is_empty(self):
        output, doc = self.generate(sign_type="internal", issuing_number="",
                                    blank_document_number=True, blank_number_line_pt=19,
                                    copy_to="测试范围", issue_unit="测试单位")
        self.assertTrue(output.name.startswith("〔发文字号待补〕："))
        self.assertEqual(doc.paragraphs[2].text, "")
        self.assertEqual(doc.paragraphs[2].paragraph_format.line_spacing, Pt(19))
        result = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_red_header_docx.py")),
                                 str(output), "--title", TITLE, "--sign-type", "internal",
                                 "--blank-document-number", "--blank-number-line-pt", "19", "--json"],
                                check=True, capture_output=True, text=True)
        checks = json.loads(result.stdout)["checks"]
        self.assertTrue(any(c["rule"] == "文号正文按要求留空且保留设定行距"
                            and c["status"] == "pass" for c in checks))
        self.assertFalse(any(c["rule"] == "文号宋体中文、Times New Roman数字、四号不加粗"
                             for c in checks))

    def test_blank_number_default_keeps_standard_line_height(self):
        _, doc = self.generate(sign_type="external", issuing_number="", blank_document_number=True)
        self.assertEqual(doc.paragraphs[2].text, "")
        self.assertEqual(doc.paragraphs[2].paragraph_format.line_spacing, Pt(27))

    def custom_source(self, body=None, no_recipient=False, table=False):
        doc = Document()
        doc.styles['Normal'].paragraph_format.space_before = Pt(0)
        doc.styles['Normal'].paragraph_format.space_after = Pt(0)
        doc.add_paragraph(TITLE)
        if not no_recipient:
            doc.add_paragraph('公司各部门：')
        for text in body or ['一、工作要求', '（一）按既有要求办理。']:
            doc.add_paragraph(text)
        if table:
            tbl = doc.add_table(rows=2, cols=2)
            tbl.cell(0, 0).text = '项目'
            tbl.cell(0, 1).text = '说明'
            tbl.cell(1, 0).text = '测试'
            tbl.cell(1, 1).text = '保留原表内容'
            borders = OxmlElement('w:tblBorders')
            for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                border = OxmlElement('w:' + name)
                border.set(qn('w:val'), 'single')
                borders.append(border)
            tbl._tbl.tblPr.append(borders)
        doc.add_paragraph(COMPANY)
        doc.add_paragraph('2026年9月1日')
        path = self.root / 'custom.docx'
        doc.save(path)
        return path

    def apply_custom(self, source, **kwargs):
        return apply(source, self.root / 'custom-output/out.docx', title=TITLE, issuing_number=NUMBER, **kwargs)

    def validate(self, output, *extra):
        result = subprocess.run([sys.executable, str(Path(__file__).with_name('validate_red_header_docx.py')),
                                 str(output), '--title', TITLE, '--json', *extra], capture_output=True, text=True)
        self.assertFalse(result.stderr, result.stderr)
        return result.returncode, json.loads(result.stdout)

    def failures(self, report):
        return [c for c in report['checks'] if c['status'] == 'fail']

    def test_heading_and_narrative_have_distinct_weights(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        heading = next(p for p in doc.paragraphs if p.text == '一、工作要求')
        body = next(p for p in doc.paragraphs if p.text.startswith('（一）'))
        self.assertTrue(all(r.bold for r in heading.runs if r.text.strip()))
        self.assertTrue(all(r.bold is False for r in body.runs if r.text.strip()))
        code, report = self.validate(output)
        self.assertEqual(code, 0, self.failures(report))
        heading.runs[0].bold = False
        doc.save(output)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '一级独立小标题仿宋四号且加粗' for c in self.failures(report)))

    def test_no_recipient_does_not_reclassify_first_body(self):
        output = self.apply_custom(self.custom_source(no_recipient=True), no_recipient=True)
        code, report = self.validate(output, '--no-recipient')
        self.assertEqual(code, 0, self.failures(report))
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '主送单位可明确识别' for c in self.failures(report)))

    def test_hierarchy_skip_is_a_hard_failure(self):
        output = self.apply_custom(self.custom_source(['一、工作要求', '（一）具体如下：', '（1）登记完成情况。']))
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '显式序号层级无跳级' for c in self.failures(report)))
        self.assertGreater(report['summary']['failed'], 0)
        self.assertFalse(report['summary']['automatic_ok'])

    def test_valid_hierarchy_and_intro_colon_are_accepted(self):
        output = self.apply_custom(self.custom_source(['一、工作要求', '（一）具体如下：', '1.按既有要求办理。', '（1）登记完成情况。']))
        code, report = self.validate(output)
        self.assertEqual(code, 0, self.failures(report))

    def test_book_title_separator_and_missing_period_are_detected(self):
        output = self.apply_custom(self.custom_source(['一、工作要求', '（一）对照《甲》、《乙》执行；']))
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        rules = [c['rule'] for c in self.failures(report)]
        self.assertIn('相邻书名号之间无顿号', rules)
        self.assertIn('完整序号条目用句号结束', rules)

    def test_closing_and_attachment_separators_are_normalized(self):
        output = self.apply_custom(self.custom_source(['一、工作要求', '（一）按既有要求办理。', '特此通知', '附件：1.测试清单', '2.测试说明']))
        code, report = self.validate(output)
        self.assertEqual(code, 0, self.failures(report))
        doc = Document(output)
        attachment = next(i for i, p in enumerate(doc.paragraphs) if p.text.startswith('附件'))
        self.assertEqual(doc.paragraphs[attachment - 1].text, '')
        doc.paragraphs[attachment - 1]._p.getparent().remove(doc.paragraphs[attachment - 1]._p)
        doc.save(output)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any('恰好空一行' in c['rule'] for c in self.failures(report)))

    def test_closing_period_is_reported_without_rewriting_text(self):
        output = self.apply_custom(self.custom_source(['一、工作要求', '（一）按既有要求办理。', '特此通知。']))
        self.assertIn('特此通知。', [p.text for p in Document(output).paragraphs])
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'].startswith('已识别结束语不带结束标点') for c in self.failures(report)))

    def test_inherited_header_indent_is_detected(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        ind = doc.styles['红头公文基准']._element.pPr.find(qn('w:ind'))
        ind.set(qn('w:firstLineChars'), '200')
        doc.save(output)
        self.assertFalse(doc.paragraphs[0]._p.xpath('./w:pPr/w:ind'))
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any('红头有效缩进' in c['rule'] for c in self.failures(report)))

    def test_character_unit_spacing_and_grid_cannot_pass_as_zero(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        paragraph = next(p for p in doc.paragraphs if p.text.startswith('（一）'))
        paragraph._p.pPr.find(qn('w:spacing')).set(qn('w:beforeLines'), '50')
        paragraph._p.pPr.find(qn('w:snapToGrid')).set(qn('w:val'), '1')
        doc.save(output)
        self.assertEqual(paragraph.paragraph_format.space_before, Pt(0))
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any('正文段前段后' in c['rule'] for c in self.failures(report)))

    def test_title_inherited_spacing_is_checked(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        doc.styles['Normal'].paragraph_format.space_after = Pt(8)
        doc.save(output)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any('标题继承后的' in c['rule'] for c in self.failures(report)))

    def test_complex_script_weight_conflict_is_detected(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        paragraph = next(p for p in doc.paragraphs if p.text.startswith('（一）'))
        paragraph.runs[0]._r.rPr.find(qn('w:bCs')).set(qn('w:val'), '1')
        doc.save(output)
        self.assertIs(paragraph.runs[0].bold, False)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '普通正文仿宋四号，保留原稿加粗' for c in self.failures(report)))

    def test_lowercase_complex_script_theme_is_detected(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        body = next(p for p in doc.paragraphs if p.text.startswith('（一）'))
        body.runs[0]._r.get_or_add_rPr().rFonts.set(qn('w:cstheme'), 'minorBidi')
        doc.save(output)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any('theme_font' in str(c) for c in self.failures(report)))

    def test_leading_zero_and_invalid_date_are_detected(self):
        output = self.apply_custom(self.custom_source())
        for value in ('2026年09月01日', '2026年2月30日'):
            doc = Document(output)
            p = next(p for p in doc.paragraphs if '年' in p.text and p.text.endswith('日'))
            p.runs[0].text = value
            doc.save(output)
            code, report = self.validate(output)
            self.assertEqual(code, 1)
            self.assertTrue(any('落款日期为有效' in c['rule'] for c in self.failures(report)))

    def test_company_and_mine_cannot_share_paragraph(self):
        source = self.custom_source()
        doc = Document(source)
        p = next(p for p in doc.paragraphs if p.text == COMPANY)
        p.add_run('大饭铺煤矿')
        doc.save(source)
        output = self.apply_custom(source)
        code, report = self.validate(output, '--expected-mine', '大饭铺煤矿')
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '公司与矿名在独立段落分行' for c in self.failures(report)))

    def test_table_xml_and_source_styles_are_preserved(self):
        source = self.custom_source(table=True)
        doc = Document(source)
        doc.styles['Normal'].paragraph_format.first_line_indent = Pt(28)
        doc.save(source)
        table_before = Document(source).tables[0]._tbl.xml
        normal_before = Document(source).styles['Normal']._element.xml
        source_before = source.read_bytes()
        output = self.apply_custom(source)
        self.assertEqual(source_before, source.read_bytes())
        generated = Document(output)
        self.assertEqual(generated.tables[0]._tbl.xml, table_before)
        self.assertEqual(generated.styles['Normal']._element.xml, normal_before)
        code, report = self.validate(output)
        self.assertEqual(code, 0, self.failures(report))

    def test_nonzero_normal_spacing_with_table_stops_without_output(self):
        source = self.custom_source(table=True)
        doc = Document(source)
        doc.styles['Normal'].paragraph_format.space_after = Pt(8)
        doc.save(source)
        before = source.read_bytes()
        with self.assertRaisesRegex(ValueError, 'Normal样式继承'):
            self.apply_custom(source)
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((self.root / 'custom-output').exists())

    def test_conflicting_public_style_is_not_rewritten(self):
        source = self.custom_source(table=True)
        doc = Document(source)
        from docx.enum.style import WD_STYLE_TYPE
        style = doc.styles.add_style('红头公文基准', WD_STYLE_TYPE.PARAGRAPH)
        style.paragraph_format.space_after = Pt(8)
        doc.tables[0].cell(0, 0).paragraphs[0].style = style
        doc.save(source)
        before = source.read_bytes()
        with self.assertRaisesRegex(ValueError, '不能全局改写同名样式'):
            self.apply_custom(source)
        self.assertEqual(source.read_bytes(), before)
        self.assertFalse((self.root / 'custom-output').exists())

    def test_reapplying_canonical_public_style_preserves_table(self):
        output = self.apply_custom(self.custom_source(table=True))
        doc = Document(output)
        table_before = doc.tables[0]._tbl.xml
        style_before = doc.styles['红头公文基准']._element.xml
        second = apply(output, self.root / 'second-output/out.docx', title=TITLE, issuing_number=NUMBER)
        generated = Document(second)
        self.assertEqual(generated.tables[0]._tbl.xml, table_before)
        self.assertEqual(generated.styles['红头公文基准']._element.xml, style_before)
        code, report = self.validate(second)
        self.assertEqual(code, 0, self.failures(report))

    def test_comments_keep_precise_anchors_during_date_formatting(self):
        source = self.custom_source()
        doc = Document(source)
        company = next(p for p in doc.paragraphs if p.text == COMPANY)
        date = doc.paragraphs[-1]
        date.runs[0].text = '2026 年 09 月 01 日'
        doc.add_comment(company.runs, text='落款前空一行', author='测试审阅人')
        doc.add_comment(date.runs, text='日期不加前导0', author='测试审阅人')
        doc.save(source)
        original = source.read_bytes()
        output = self.apply_custom(source)
        record = extract(output)
        self.assertEqual(record['comment_count'], 2)
        by_text = {c['text']: c for c in record['comments']}
        self.assertEqual(by_text['日期不加前导0']['anchors'][0]['selected_text'], '2026年9月1日')
        self.assertEqual(by_text['落款前空一行']['anchors'][0]['selected_text'], COMPANY)
        self.assertEqual(source.read_bytes(), original)

    def test_comment_on_old_header_prevents_silent_anchor_deletion(self):
        source = self.source(decorated=True)
        doc = Document(source)
        doc.add_comment(doc.paragraphs[0].runs, text='核对红头', author='测试审阅人')
        doc.save(source)
        before = source.read_bytes()
        with self.assertRaisesRegex(ValueError, '原红头包含批注'):
            self.apply_custom(source)
        self.assertEqual(source.read_bytes(), before)

    def test_bookmarked_empty_paragraph_is_preserved(self):
        source = self.custom_source()
        doc = Document(source)
        company = next(p for p in doc.paragraphs if p.text == COMPANY)
        blank = company.insert_paragraph_before()
        start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
        start.set(qn('w:id'), '42'); start.set(qn('w:name'), 'test_anchor')
        end.set(qn('w:id'), '42')
        blank._p.extend([start, end])
        doc.save(source)
        output = self.apply_custom(source)
        self.assertTrue(Document(output)._element.body.xpath('.//w:bookmarkStart[@w:id="42"]'))
        _, report = self.validate(output)
        self.assertTrue(any(c['rule'] == '带标记或对象空段的实际分隔间距' and c['status'] == 'manual' for c in report['checks']))

    def test_manual_pending_does_not_mean_formal_ready(self):
        output = self.apply_custom(self.custom_source())
        code, report = self.validate(output, '--require-manual-review')
        self.assertEqual(code, 2)
        self.assertEqual(report['summary']['failed'], 0)
        self.assertTrue(report['summary']['automatic_ok'])
        self.assertFalse(report['summary']['formal_ready'])
        self.assertGreater(report['summary']['pending_manual'], 0)

    def test_partial_source_bold_remains_partial(self):
        source = self.custom_source(body=['本条按原内容办理。'])
        doc = Document(source)
        p = doc.paragraphs[2]
        p.clear()
        p.add_run('本条').bold = True
        p.add_run('按原内容办理。').bold = False
        doc.add_comment(p.runs[0], text='原有强调', author='测试审阅人')
        doc.save(source)
        before = source.read_bytes()
        output = self.apply_custom(source)
        body = next(p for p in Document(output).paragraphs if p.text == '本条按原内容办理。')
        self.assertEqual(bold_snapshot(body)[1], [True, True] + [False] * len('按原内容办理。'))
        self.assertEqual(extract(output)['comments'][0]['anchors'][0]['selected_text'], '本条')
        self.assertEqual(before, source.read_bytes())
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_source_style_bold_is_materialized(self):
        from docx.enum.style import WD_STYLE_TYPE
        source = self.custom_source(body=['原样保留本条。'])
        doc = Document(source)
        style = doc.styles.add_style('测试原稿强调', WD_STYLE_TYPE.PARAGRAPH)
        style.font.bold = True
        doc.paragraphs[2].style = style
        self.assertIsNone(doc.paragraphs[2].runs[0].bold)
        doc.save(source)
        output = self.apply_custom(source)
        body = next(p for p in Document(output).paragraphs if p.text == '原样保留本条。')
        self.assertTrue(all(r.bold for r in body.runs))
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_style_bold_toggles_are_not_mistaken_for_source_bold(self):
        from docx.enum.style import WD_STYLE_TYPE
        source = self.custom_source(body=['原样保留本条。'])
        doc = Document(source)
        ps = doc.styles.add_style('测试段落强调', WD_STYLE_TYPE.PARAGRAPH)
        ps.font.bold = True
        rs = doc.styles.add_style('测试字符切换', WD_STYLE_TYPE.CHARACTER)
        rs.font.bold = True
        doc.paragraphs[2].style = ps
        doc.paragraphs[2].runs[0].style = rs
        self.assertFalse(effective_bold(doc.paragraphs[2].runs[0]))
        doc.save(source)
        output = self.apply_custom(source)
        body = next(p for p in Document(output).paragraphs if p.text == '原样保留本条。')
        self.assertFalse(effective_bold(body.runs[0]))

    def test_title_reflow_preserves_partial_source_bold(self):
        source = self.custom_source()
        doc = Document(source)
        p = doc.paragraphs[0]
        p.clear()
        p.add_run('关于').bold = False
        p.add_run('测试事项').bold = True
        p.add_run('的通知').bold = False
        doc.save(source)
        output = self.apply_custom(source, title_lines=['关于测试事项', '的通知'])
        p = next(p for p in Document(output).paragraphs if p.text.replace('\n', '') == TITLE)
        value, weights, _ = compact_snapshot(bold_snapshot(p))
        self.assertEqual(value, TITLE)
        self.assertEqual(weights, [False] * 2 + [True] * 4 + [False] * 3)
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_signature_date_normalization_preserves_emphasis(self):
        source = self.custom_source()
        doc = Document(source)
        doc.paragraphs[-1].runs[0].text = '2026 年 09 月 01 日'
        doc.paragraphs[-1].runs[0].bold = True
        doc.save(source)
        output = self.apply_custom(source)
        p = next(p for p in Document(output).paragraphs if p.text == '2026年9月1日')
        self.assertTrue(all(effective_bold(r) for r in p.runs if r.text.strip()))
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_validator_detects_erased_source_bold(self):
        source = self.custom_source(body=['本条继续保留。'])
        doc = Document(source)
        doc.paragraphs[2].runs[0].bold = True
        doc.save(source)
        output = self.apply_custom(source)
        doc = Document(output)
        p = next(p for p in doc.paragraphs if p.text == '本条继续保留。')
        for r in p.runs:
            r.bold = False
            r._r.rPr.find(qn('w:bCs')).set(qn('w:val'), '0')
        doc.save(output)
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '原稿加粗文字范围保留' for c in self.failures(report)))

    def test_validator_detects_extended_partial_source_bold(self):
        source = self.custom_source(body=['本条继续保留。'])
        doc = Document(source)
        p = doc.paragraphs[2]
        p.clear()
        p.add_run('本条').bold = True
        p.add_run('继续保留。').bold = False
        doc.save(source)
        output = self.apply_custom(source)
        doc = Document(output)
        p = next(p for p in doc.paragraphs if p.text == '本条继续保留。')
        for r in p.runs:
            r.bold = True
            r._r.rPr.find(qn('w:bCs')).set(qn('w:val'), '1')
        doc.save(output)
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 1)

    def test_trailing_empty_paragraphs_and_page_breaks_are_removed(self):
        source = self.custom_source()
        doc = Document(source)
        doc.paragraphs[-1].runs[0].add_break(__import__('docx.enum.text', fromlist=['WD_BREAK']).WD_BREAK.PAGE)
        for _ in range(30):
            doc.add_paragraph()
        doc.add_page_break()
        doc.save(source)
        before = source.read_bytes()
        output = self.apply_custom(source)
        doc = Document(output)
        self.assertEqual(doc.paragraphs[-1].text, '2026年9月1日')
        self.assertFalse(doc.paragraphs[-1]._p.xpath('.//w:br'))
        self.assertEqual(before, source.read_bytes())
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_empty_final_section_keeps_real_content_section_settings(self):
        from docx.enum.section import WD_SECTION
        from docx.shared import Mm
        source = self.custom_source()
        doc = Document(source)
        doc.sections[-1].right_margin = Mm(23)
        doc.add_section(WD_SECTION.ODD_PAGE).right_margin = Mm(40)
        doc.add_paragraph()
        doc.save(source)
        output = self.apply_custom(source)
        doc = Document(output)
        self.assertEqual(len(doc.sections), 1)
        self.assertLess(abs(doc.sections[0].right_margin - Mm(23)), Mm(0.01))
        self.assertEqual(doc.paragraphs[-1].text, '2026年9月1日')

    def test_real_attachment_after_signature_is_not_deleted(self):
        source = self.custom_source()
        doc = Document(source)
        p = doc.add_paragraph('附件：另附业务材料。')
        p.runs[0].bold = True
        doc.add_paragraph('原附件内容继续保留。')
        doc.save(source)
        output = self.apply_custom(source)
        texts = [p.text for p in Document(output).paragraphs]
        self.assertIn('附件：另附业务材料。', texts)
        self.assertEqual(texts[-1], '原附件内容继续保留。')

    def test_validator_rejects_trailing_blank_suffix(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        doc.add_page_break()
        doc.save(output)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '文末无可安全识别的空段及末尾分页' for c in self.failures(report)))

    def test_appointment_notice_keeps_original_arabic_items(self):
        source = self.custom_source(body=['1.按原内容办理。', '2.原有职责继续保留。'])
        output = self.apply_custom(source, inspection_material=True, appointment_notice=True)
        doc = Document(output)
        texts = [p.text for p in doc.paragraphs]
        self.assertIn('1.按原内容办理。', texts)
        self.assertFalse(any(p.text.startswith('一、') for p in doc.paragraphs))
        code, report = self.validate(output, '--source', str(source), '--inspection-material', '--appointment-notice')
        self.assertEqual(code, 0, self.failures(report))

    def test_titleless_appointment_notice_adds_no_title(self):
        source = self.custom_source(body=['1.按原内容办理。'], no_recipient=True)
        doc = Document(source)
        doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p)
        doc.save(source)
        output = self.apply_custom(source, inspection_material=True, appointment_notice=True, no_title=True)
        doc = Document(output)
        self.assertFalse(any(p.text == TITLE for p in doc.paragraphs))
        self.assertIn('1.按原内容办理。', [p.text for p in doc.paragraphs])
        code, report = self.validate(output, '--source', str(source), '--inspection-material', '--appointment-notice', '--no-title')
        self.assertEqual(code, 0, self.failures(report))

    def test_appointment_notice_keeps_original_punctuation(self):
        body = ['1.按原内容办理；', '2.参照《示例规范甲》、《示例规范乙》。', '特此通知。']
        source = self.custom_source(body=body)
        output = self.apply_custom(source, inspection_material=True, appointment_notice=True)
        texts = [p.text for p in Document(output).paragraphs]
        self.assertTrue(all(text in texts for text in body))
        code, report = self.validate(output, '--source', str(source), '--inspection-material', '--appointment-notice')
        self.assertEqual(code, 0, self.failures(report))

    def test_appointment_notice_does_not_force_heading_bold(self):
        source = self.custom_source(body=['一、原有独立说明', '原有内容继续保留。'])
        output = self.apply_custom(source, inspection_material=True, appointment_notice=True)
        heading = next(p for p in Document(output).paragraphs if p.text == '一、原有独立说明')
        self.assertFalse(any(effective_bold(r) for r in heading.runs))
        code, report = self.validate(output, '--source', str(source), '--inspection-material', '--appointment-notice')
        self.assertEqual(code, 0, self.failures(report))

    def test_appointment_validator_detects_invented_heading(self):
        source = self.custom_source(body=['按原内容办理。'])
        output = self.apply_custom(source, inspection_material=True, appointment_notice=True)
        doc = Document(output)
        next(p for p in doc.paragraphs if p.text == '按原内容办理。').runs[0].text = '一、按原内容办理。'
        doc.save(output)
        code, report = self.validate(output, '--source', str(source), '--inspection-material', '--appointment-notice')
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '任职迎检通知原内容未增写标题、序号或改写条目' for c in self.failures(report)))

    def test_appointment_mode_requires_inspection_scope(self):
        with self.assertRaisesRegex(ValueError, '仅用于已确认的迎检资料'):
            self.apply_custom(self.custom_source(), appointment_notice=True)
        with self.assertRaisesRegex(ValueError, '须同时声明任职类迎检通知'):
            self.apply_custom(self.custom_source(), no_title=True)

    def test_web_contract_keeps_appointment_and_no_title_scope(self):
        path = self.root / 'task.json'
        task = {'task_type': 'office_red_head_document', 'skill_name': 'red-header-official-document',
                'output_folder': str(self.root / 'output'), 'payload': {'body': '按原内容办理。',
                'title': '原稿内容标签', 'inspectionMaterial': True, 'appointmentNotice': True, 'noTitle': True}}
        script = str(Path(__file__).with_name('validate_task.py'))
        for inspection, expected in [(True, 0), (False, 1)]:
            task['payload']['inspectionMaterial'] = inspection
            path.write_text(json.dumps(task, ensure_ascii=False))
            result = subprocess.run([sys.executable, script, str(path)], capture_output=True, text=True)
            self.assertEqual(result.returncode, expected)
            report = json.loads(result.stdout)
            self.assertIs(report['appointment_notice'], True)
            self.assertIs(report['no_title'], True)
        task['payload']['inspectionMaterial'] = True
        task['payload']['noTitle'] = 'true'
        path.write_text(json.dumps(task, ensure_ascii=False))
        result = subprocess.run([sys.executable, script, str(path)], capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)

    def test_hidden_work_and_wps_lock_files_are_not_deliverables(self):
        folder = self.root / 'delivery'; folder.mkdir()
        work = folder / '.work'; work.mkdir()
        (work / 'source.docx').write_bytes(b'test')
        (work / 'preview.pdf').write_bytes(b'test')
        (folder / '~$lock.docx').write_bytes(b'test')
        self.assertEqual(visible_files(folder), [])
        (folder / 'deliverable.docx').write_bytes(b'test')
        self.assertEqual([p.name for p in visible_files(folder)], ['deliverable.docx'])


    def test_existing_source_number_is_used_when_argument_is_omitted(self):
        source = self.source(decorated=True)
        original = source.read_bytes()
        output = apply(source, self.root / 'auto/out.docx', title=TITLE, sign_type='internal', inspection_material=True)
        self.assertEqual(output.name, NUMBER + '：' + TITLE + '.docx')
        self.assertEqual(Document(output).paragraphs[2].text, NUMBER)
        self.assertEqual(source.read_bytes(), original)

    def test_confirmed_number_overrides_existing_source_number(self):
        number = '蒙力煤字〔2026〕2号'
        output = apply(self.source(decorated=True), self.root / 'override/out.docx', title=TITLE,
                       issuing_number=number, inspection_material=True)
        self.assertEqual(output.name, number + '：' + TITLE + '.docx')
        self.assertNotIn(NUMBER, [p.text for p in Document(output).paragraphs])

    def test_ambiguous_source_numbers_are_not_chosen_automatically(self):
        source = self.source(decorated=True)
        doc = Document(source)
        doc.paragraphs[2].insert_paragraph_before('蒙力煤字〔2026〕2号')
        doc.save(source)
        with self.assertRaisesRegex(ValueError, '多个不同文号'):
            apply(source, self.root / 'ambiguous/out.docx', title=TITLE, inspection_material=True)
        self.assertFalse((self.root / 'ambiguous').exists())

    def test_missing_source_number_stays_pending(self):
        output = apply(self.source(), self.root / 'pending/out.docx', title=TITLE)
        self.assertTrue(output.name.startswith('〔发文字号待补〕：'))
        self.assertEqual(Document(output).paragraphs[2].text, '〔发文字号待补〕')

    def test_explicit_blank_number_does_not_reuse_source_number(self):
        output = apply(self.source(decorated=True), self.root / 'blank/out.docx', title=TITLE,
                       inspection_material=True, blank_document_number=True)
        self.assertTrue(output.name.startswith('〔发文字号待补〕：'))
        self.assertEqual(Document(output).paragraphs[2].text, '')

    def test_centered_body_is_detected_directly_and_through_style(self):
        for inherited in (False, True):
            with self.subTest(inherited=inherited):
                output = self.apply_custom(self.custom_source(body=['1.按原内容办理。']),
                                           inspection_material=True, appointment_notice=True)
                doc = Document(output)
                p = next(p for p in doc.paragraphs if p.text.startswith('1.'))
                p.alignment = None if inherited else WD_ALIGN_PARAGRAPH.CENTER
                if inherited:
                    p.style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.save(output)
                code, report = self.validate(output, '--inspection-material', '--appointment-notice')
                self.assertEqual(code, 1)
                self.assertTrue(any(c['rule'] == '主送及普通正文左对齐，任职条目不居中' for c in self.failures(report)))

    def source_with_mine(self):
        source = self.custom_source()
        doc = Document(source)
        doc.paragraphs[-1].insert_paragraph_before('大饭铺煤矿')
        doc.save(source)
        return source

    def test_mine_must_share_signature_alignment_region(self):
        for wrong_alignment in (True, False):
            with self.subTest(wrong_alignment=wrong_alignment):
                output = self.apply_custom(self.source_with_mine())
                doc = Document(output)
                p = next(p for p in doc.paragraphs if p.text == '大饭铺煤矿')
                if wrong_alignment:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.paragraph_format.left_indent = Pt(0)
                    p._p.pPr.find(qn('w:ind')).set(qn('w:start'), '0')
                doc.save(output)
                code, report = self.validate(output)
                self.assertEqual(code, 1)
                self.assertTrue(any(c['rule'] == '公司落款、矿名与日期在同一对齐区域居中' for c in self.failures(report)))

    def test_signature_first_line_character_indent_is_detected(self):
        output = self.apply_custom(self.custom_source())
        doc = Document(output)
        p = doc.paragraphs[-1]
        p._p.pPr.find(qn('w:ind')).set(qn('w:firstLineChars'), '2100')
        doc.save(output)
        code, report = self.validate(output)
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'].startswith('落款无首行/悬挂字符缩进') for c in self.failures(report)))

    def test_signature_padding_and_tabs_are_detected(self):
        for padding in (' ' * 12, '\t'):
            with self.subTest(padding=repr(padding)):
                output = self.apply_custom(self.source_with_mine())
                doc = Document(output)
                p = next(p for p in doc.paragraphs if p.text == '大饭铺煤矿')
                p.runs[0].text = padding + p.runs[0].text
                doc.save(output)
                code, report = self.validate(output)
                self.assertEqual(code, 1)
                self.assertTrue(any(c['rule'].startswith('落款无首行/悬挂字符缩进') for c in self.failures(report)))

    def test_filename_number_is_checked_without_title_argument(self):
        source = self.custom_source(body=['按原内容办理。'], no_recipient=True)
        doc = Document(source)
        doc.paragraphs[0]._p.getparent().remove(doc.paragraphs[0]._p)
        doc.save(source)
        output = self.apply_custom(source, inspection_material=True, appointment_notice=True, no_title=True)
        command = [sys.executable, str(Path(__file__).with_name('validate_red_header_docx.py')),
                   str(output), '--inspection-material', '--appointment-notice', '--no-title', '--no-recipient', '--json']
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        wrong = output.with_name('原稿内容标签.docx')
        output.rename(wrong)
        command[2] = str(wrong)
        result = subprocess.run(command, capture_output=True, text=True)
        self.assertEqual(result.returncode, 1)
        self.assertTrue(any(c['rule'] == '未传标题时文件名仍须包含正文文号及标题标签'
                            for c in self.failures(json.loads(result.stdout))))


    def test_external_normalizes_number_in_body_and_filename(self):
        output, doc = self.generate(sign_type='external')
        expected = '蒙力煤发〔2026〕1号'
        self.assertEqual(doc.paragraphs[2].text, expected)
        self.assertEqual(output.name, expected + '：' + TITLE + '.docx')
        code, report = self.validate(output, '--sign-type', 'external', '--issuing-number', NUMBER)
        self.assertEqual(code, 0, self.failures(report))

    def test_external_source_number_is_normalized_without_argument(self):
        source = self.custom_source()
        doc = Document(source)
        doc.paragraphs[0].insert_paragraph_before(NUMBER)
        doc.save(source)
        original = source.read_bytes()
        output = apply(source, self.root / 'source-number/out.docx', title=TITLE, sign_type='external')
        self.assertEqual(Document(output).paragraphs[2].text, '蒙力煤发〔2026〕1号')
        self.assertEqual(source.read_bytes(), original)
        code, report = self.validate(output, '--source', str(source), '--sign-type', 'external')
        self.assertEqual(code, 0, self.failures(report))

    def test_external_keeps_number_year_and_serial(self):
        output, doc = self.generate(sign_type='external', issuing_number='蒙力煤字〔2025〕007号')
        expected = '蒙力煤发〔2025〕007号'
        self.assertEqual(doc.paragraphs[2].text, expected)
        self.assertEqual(output.name, expected + '：' + TITLE + '.docx')
        self.assertIn('2026年9月1日', [p.text for p in doc.paragraphs])

    def test_external_validator_rejects_internal_prefix(self):
        output, doc = self.generate(sign_type='external')
        doc.paragraphs[2].runs[0].text = NUMBER
        wrong = output.with_name(NUMBER + '：' + TITLE + '.docx')
        doc.save(wrong)
        code, report = self.validate(wrong, '--sign-type', 'external')
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'].startswith('外发文号使用蒙力煤发') for c in self.failures(report)))

    def test_web_external_number_resolution(self):
        path = self.root / 'external-task.json'
        path.write_text(json.dumps({'task_type': 'office_red_head_document', 'skill_name': 'red-header-official-document',
                                    'output_folder': str(self.root / 'output'), 'payload': {'title': TITLE,
                                    'body': '按测试要求办理。', 'recipient': '外部测试单位：', 'signType': 'external',
                                    'documentNo': '蒙力煤字〔2025〕007号'}}, ensure_ascii=False))
        result = subprocess.run([sys.executable, str(Path(__file__).with_name('validate_task.py')), str(path)],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertEqual(json.loads(result.stdout)['resolved_document_no'], '蒙力煤发〔2025〕007号')


    def test_recipient_padding_is_removed_at_each_line_start(self):
        source = self.custom_source()
        doc = Document(source)
        p = doc.paragraphs[1]
        p.clear()
        p.add_run(' \u3000')
        run = p.add_run('\t公司各部门：\n\u3000\u3000相关单位：')
        doc.add_comment(run, text='核对主送范围', author='示例审阅人')
        p.paragraph_format.first_line_indent = Pt(28)
        p.style.paragraph_format.left_indent = Pt(28)
        doc.save(source)
        before = source.read_bytes()
        output = self.apply_custom(source)
        p = next(p for p in Document(output).paragraphs if p.text.startswith('公司各部门'))
        self.assertEqual(p.text, '公司各部门：\n相关单位：')
        self.assertFalse(recipient_layout_issues(p))
        self.assertEqual(extract(output)['comments'][0]['text'], '核对主送范围')
        self.assertEqual(source.read_bytes(), before)
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_recipient_leading_padding_is_a_separate_failure(self):
        for padding in ('  ', '\u3000\u3000', '\t'):
            with self.subTest(padding=repr(padding)):
                output = self.apply_custom(self.custom_source())
                doc = Document(output)
                p = next(p for p in doc.paragraphs if p.text == '公司各部门：')
                p.runs[0].text = padding + p.runs[0].text
                doc.save(output)
                code, report = self.validate(output)
                self.assertEqual(code, 1)
                self.assertTrue(any(c['rule'].startswith('主送单位顶格：') for c in self.failures(report)))

    def test_recipient_two_character_indent_is_rejected(self):
        for inherited in (False, True):
            with self.subTest(inherited=inherited):
                output = self.apply_custom(self.custom_source())
                doc = Document(output)
                p = next(p for p in doc.paragraphs if p.text == '公司各部门：')
                layer = p.style._element.get_or_add_pPr() if inherited else p._p.pPr
                if inherited:
                    p._p.pPr.remove(p._p.pPr.find(qn('w:ind')))
                ind = layer.find(qn('w:ind'))
                if ind is None:
                    ind = OxmlElement('w:ind'); layer.append(ind)
                ind.set(qn('w:firstLineChars'), '200')
                doc.save(output)
                code, report = self.validate(output)
                self.assertEqual(code, 1)
                self.assertTrue(any(c['rule'].startswith('主送单位顶格：') for c in self.failures(report)))

    def test_blank_suffix_after_real_attachment_is_cleaned(self):
        source = self.custom_source()
        doc = Document(source)
        doc.add_paragraph('附件：示例说明')
        doc.add_paragraph('附件中的真实业务内容必须保留。')
        for _ in range(30):
            doc.add_paragraph('')
        doc.add_page_break()
        doc.save(source)
        output = self.apply_custom(source)
        result = Document(output)
        self.assertEqual(result.paragraphs[-1].text, '附件中的真实业务内容必须保留。')
        self.assertIn('附件：示例说明', [p.text for p in result.paragraphs])
        code, report = self.validate(output, '--source', str(source))
        self.assertEqual(code, 0, self.failures(report))

    def test_rendered_blank_middle_page_with_header_and_footer_is_rejected(self):
        from PIL import Image, ImageDraw
        output = self.apply_custom(self.custom_source())
        pages = self.root / 'render'; pages.mkdir()
        for n in (1, 2, 3):
            img = Image.new('RGB', (840, 1188), 'white')
            draw = ImageDraw.Draw(img)
            draw.rectangle((110, 18, 230, 25), fill='black')
            draw.rectangle((410, 1145, 430, 1152), fill='black')
            if n != 2:
                draw.rectangle((110, 145, 300, 160), fill='black')
            img.save(pages / f'page-{n}.png')
        evidence = rendered_page_content(pages, Document(output).sections[0])
        self.assertEqual([p['page'] for p in evidence if p['blank_body']], [2])
        code, report = self.validate(output, '--page-images', str(pages))
        self.assertEqual(code, 1)
        self.assertTrue(any(c['rule'] == '完整渲染页面无多余空白页' for c in self.failures(report)))
        (pages / 'page-2.png').unlink()
        code, report = self.validate(output, '--page-images', str(pages))
        self.assertEqual(code, 1)

    def test_rendered_content_page_passes_and_missing_images_remain_pending(self):
        from PIL import Image, ImageDraw
        output = self.apply_custom(self.custom_source())
        pages = self.root / 'render'; pages.mkdir()
        img = Image.new('RGB', (840, 1188), 'white')
        ImageDraw.Draw(img).rectangle((110, 145, 300, 160), fill='red')
        img.save(pages / 'page-1.png')
        code, report = self.validate(output, '--page-images', str(pages))
        self.assertEqual(code, 0, self.failures(report))
        self.assertTrue(any(c['rule'] == '完整渲染页面无多余空白页' and c['status'] == 'pass' for c in report['checks']))
        code, report = self.validate(output)
        self.assertTrue(any(c['rule'] == '完整渲染页面无多余空白页' and c['status'] == 'manual' for c in report['checks']))
        empty = self.root / 'empty'; empty.mkdir()
        code, report = self.validate(output, '--page-images', str(empty))
        self.assertEqual(code, 1)


    def execution_source(self):
        source = self.custom_source(body=['三、工作要求', '（八）示例说明：', '1、按原有内容办理。',
                                          '    四、原有独立标题', '1、正文引出要求：原有说明继续保留。'],
                                    no_recipient=True, table=True)
        doc = Document(source)
        doc.paragraphs[0].insert_paragraph_before(NUMBER)
        table = doc.tables[0]
        for p in table._tbl.xpath('.//w:p'):
            from docx.text.paragraph import Paragraph
            Paragraph(p, table).paragraph_format.keep_with_next = True
        doc.save(source)
        return source

    def test_execution_layout_preserves_content_and_formats_table(self):
        source = self.execution_source()
        before = source.read_bytes()
        shape = table_shape(Document(source).tables[0])
        output, changes = revise_execution(source, self.root / 'execution/out.docx', title=TITLE)
        doc = Document(output)
        self.assertEqual(source.read_bytes(), before)
        self.assertEqual(table_shape(doc.tables[0]), shape)
        self.assertIn('（八）示例说明', [p.text for p in doc.paragraphs])
        self.assertIn('1.正文引出要求：原有说明继续保留。', [p.text for p in doc.paragraphs])
        self.assertIn('四、原有独立标题', [p.text for p in doc.paragraphs])
        self.assertFalse(execution_issues(doc, TITLE))
        self.assertTrue(changes)

    def test_execution_checker_detects_table_style_and_heading_regressions(self):
        output, _ = revise_execution(self.execution_source(), self.root / 'execution/out.docx', title=TITLE)
        doc = Document(output)
        table = doc.tables[0]
        table.cell(0, 0).paragraphs[0].runs[0].bold = False
        table.cell(1, 0).paragraphs[0].runs[0].font.size = Pt(14)
        p = next(p for p in doc.paragraphs if p.text == '（八）示例说明')
        p.runs[0].text += '：'
        issues = execution_issues(doc, TITLE)
        self.assertTrue(any('表头加粗' in x['issue'] for x in issues))
        self.assertTrue(any('独立括号小标题仍带冒号' == x['issue'] for x in issues))

    def test_fixed_distribution_template_changes_only_subject_name(self):
        source = self.execution_source()
        doc = Document(source)
        doc.add_paragraph('主题：旧名称')
        doc.add_paragraph('抄送：旧范围')
        doc.add_paragraph('发文：旧单位\t共印1份')
        doc.save(source)
        template = {'copy_to': '通用测试范围', 'issue_unit': '通用测试单位', 'print_count': 3}
        output, _ = revise_execution(source, self.root / 'execution/out.docx', title=TITLE, distribution_template=template)
        doc = Document(output)
        self.assertEqual([p.text for p in doc.paragraphs[-3:]],
                         ['主题：' + TITLE, '抄送：通用测试范围', '发文：通用测试单位\t共印3份'])
        self.assertFalse(execution_issues(doc, TITLE, template))
        with self.assertRaisesRegex(ValueError, '不自动猜测'):
            revise_execution(source, self.root / 'bad/out.docx', title=TITLE, distribution_template={})

    def test_execution_rejects_html_and_accepts_extensionless_word_package(self):
        html = self.root / 'web.docx'; html.write_text('<html><body>测试网页</body></html>')
        with self.assertRaisesRegex(ValueError, '不是可编辑的Word'):
            revise_execution(html, self.root / 'html/out.docx', title=TITLE)
        source = self.execution_source()
        no_extension = source.with_suffix(''); source.rename(no_extension)
        output, _ = revise_execution(no_extension, self.root / 'word/out', title=TITLE)
        self.assertEqual(output.suffix, '.docx')
        self.assertTrue(output.name.startswith(NUMBER + '：'))
        self.assertFalse(execution_issues(Document(output), TITLE))

    def test_confirmed_small_table_grouping_has_an_explicit_scope(self):
        source = self.execution_source()
        output, _ = revise_execution(source, self.root / 'execution/out.docx', title=TITLE, keep_together_tables=[1])
        self.assertFalse(execution_issues(Document(output), TITLE, keep_together_tables=[1]))
        with self.assertRaisesRegex(ValueError, '超出原稿表格范围'):
            revise_execution(source, self.root / 'bad/out.docx', title=TITLE, keep_together_tables=[2])

    def test_authorized_number_punctuation_changes_keep_source_emphasis(self):
        from revise_internal_execution_docx import normalize_execution_paragraph
        from red_header_common import source_bold_map
        source = self.execution_source()
        doc = Document(source)
        for p in doc.paragraphs:
            if p.text.startswith(('1、', '（八）')):
                for r in p.runs:
                    r.bold = True
        doc.save(source)
        output, _ = revise_execution(source, self.root / 'execution/out.docx', title=TITLE)
        reference = Document(source)
        for p in reference.paragraphs:
            normalize_execution_paragraph(p)
        _, dropped, equal = source_bold_map(reference, Document(output))
        self.assertFalse(dropped)
        self.assertTrue(equal)


if __name__ == "__main__":
    unittest.main()
